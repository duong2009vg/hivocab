"""
Upload IELTS Vol 1 to Vol 9 vocabulary and passages to Supabase.
Handles:
- 9 Topics ('IELTS Vol 1' to 'IELTS Vol 9') under category 'IELTS Actual Tests'
- 100 Tests with proper ordering
- 300 Passages with bilingual text (content_en, content_vi) from docx
- ~57,500 deduplicated words with intra-passage deduplication
- Special cases:
    * Vol 5 - Test 8: First 180 clean rows of Passage 1
    * Vol 2 - Test 4: Supplement 195 words for Passage 3 from docx Table 5
    * Vol 9 - Test 1: Standard passage titles (no docx available)
    * Title overrides for 3 docx with omitted titles
"""

import os
import re
import sys
import uuid
import json
import requests
import openpyxl
import docx

ROOT_DIR = r"c:\Users\MY PC\Documents\GitHub\hivocab"
DATA_DIR = os.path.join(ROOT_DIR, "data", "full bộ vol")
SUPABASE_URL = "https://swehdtrqjyklmsefkjdf.supabase.co"
ANON_KEY = (
    "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9."
    "eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6InN3ZWhkdHJxanlrbG1zZWZramRmIiwicm9sZSI6ImFub24iLCJpYXQiOjE3NzgzOTc4MDcsImV4cCI6MjA5Mzk3MzgwN30."
    "dXRhEmvS8J21aJ3dwZ4jHaWuKbhNw2yys90YTIop2EU"
)

OVERRIDE_TITLES = {
    ("VOL 2", "Vol 2 - Test 8", 2): "THE LITTLE ICE AGE",
    ("VOL 3", "Vol 3 - Test 6", 2): "THE ORIGIN OF BEES",
    ("VOL 5", "Vol 5 - Test 7", 3): "EVOLUTIONARY PSYCHOLOGY AT WORK",
}

def extract_titles_and_texts(docx_path, vol_name, test_base):
    if not os.path.exists(docx_path):
        return [
            {"title": f"Passage {i}", "content_en": None, "content_vi": None}
            for i in (1, 2, 3)
        ]
    
    doc = docx.Document(docx_path)
    
    # Extract passage texts from tables:
    # Table 0 -> P1, Table 2 -> P2, Table 4 -> P3
    passage_texts = []
    text_table_indices = [0, 2, 4]
    for idx, t_idx in enumerate(text_table_indices):
        en_text = None
        vi_text = None
        if len(doc.tables) > t_idx:
            t = doc.tables[t_idx]
            if len(t.rows) >= 1 and len(t.columns) >= 2:
                en_text = t.rows[0].cells[0].text.strip() or None
                vi_text = t.rows[0].cells[1].text.strip() or None
        passage_texts.append((en_text, vi_text))
        
    # Extract titles from paragraphs
    titles = [None, None, None]
    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for p_idx, p_num in enumerate((1, 2, 3)):
        if (vol_name, test_base, p_num) in OVERRIDE_TITLES:
            titles[p_idx] = OVERRIDE_TITLES[(vol_name, test_base, p_num)]
            continue
        for i, line in enumerate(lines):
            if f"PASSAGE {p_num}" in line.upper():
                for j in range(i, min(i + 10, len(lines))):
                    l = lines[j]
                    if "DỊCH ĐỀ" in l.upper() or "DICH DE" in l.upper():
                        t = l.split(":")[-1].strip()
                        if t and t.upper() not in ["DỊCH ĐỀ", "DICH DE", "TỪ VỰNG:"]:
                            titles[p_idx] = t
                            break
                        elif j + 1 < len(lines):
                            cand = lines[j + 1].strip()
                            if cand.upper() != "TỪ VỰNG:":
                                titles[p_idx] = cand
                                break
                if titles[p_idx]:
                    break
        if not titles[p_idx] or "TỪ VỰNG" in titles[p_idx].upper():
            titles[p_idx] = f"Passage {p_num}"
            
    res = []
    for i in range(3):
        res.append({
            "title": titles[i],
            "content_en": passage_texts[i][0],
            "content_vi": passage_texts[i][1]
        })
    return res

def parse_all_data():
    vols = sorted([d for d in os.listdir(DATA_DIR) if os.path.isdir(os.path.join(DATA_DIR, d))])
    
    topics = []
    tests = []
    passages = []
    words = []
    
    print(f"Found {len(vols)} volumes to process.")
    
    for v in vols:
        vol_num = int(re.search(r"\d+", v).group(0))
        topic_name = f"IELTS Vol {vol_num}"
        topic_id = str(uuid.uuid5(uuid.NAMESPACE_DNS, f"hivocab.ielts.vol.{vol_num}"))
        
        topics.append({
            "id": topic_id,
            "name": topic_name,
            "icon": "menu_book",
            "category": "IELTS Actual Tests"
        })
        
        v_path = os.path.join(DATA_DIR, v)
        xlsx_files = [f for f in os.listdir(v_path) if (f.endswith(".xlsx") or f.endswith(".xls")) and not f.startswith("~$")]
        
        def test_sort_key(filename):
            match = re.search(r"Test\s*(\d+)", filename, re.IGNORECASE)
            return int(match.group(1)) if match else 999
            
        xlsx_files.sort(key=test_sort_key)
        
        for xf in xlsx_files:
            test_order = test_sort_key(xf)
            test_name = f"Test {test_order}"
            test_id = str(uuid.uuid5(uuid.UUID(topic_id), f"test.{test_order}"))
            test_base = os.path.splitext(xf)[0]
            
            tests.append({
                "id": test_id,
                "topic_id": topic_id,
                "name": test_name,
                "test_order": test_order
            })
            
            docx_path = os.path.join(v_path, test_base + ".docx")
            passage_meta = extract_titles_and_texts(docx_path, v, test_base)
            
            xl_path = os.path.join(v_path, xf)
            wb = openpyxl.load_workbook(xl_path, data_only=True)
            
            for p_num in (1, 2, 3):
                p_meta = passage_meta[p_num - 1]
                passage_id = str(uuid.uuid5(uuid.UUID(test_id), f"passage.{p_num}"))
                
                passages.append({
                    "id": passage_id,
                    "test_id": test_id,
                    "topic_id": topic_id,
                    "passage_number": p_num,
                    "title": p_meta["title"],
                    "topic_label": "",
                    "content_en": p_meta["content_en"],
                    "content_vi": p_meta["content_vi"]
                })
                
                words_raw = []
                if v == "VOL 2" and test_order == 4 and p_num == 3:
                    if os.path.exists(docx_path):
                        doc = docx.Document(docx_path)
                        if len(doc.tables) >= 6:
                            for row in doc.tables[5].rows:
                                w_text = row.cells[0].text.strip()
                                m_text = row.cells[1].text.strip() if len(row.cells) > 1 else ""
                                if w_text:
                                    words_raw.append((w_text, m_text))
                else:
                    sheet_name = f"PASSAGE {p_num}"
                    if sheet_name in wb.sheetnames:
                        sheet = wb[sheet_name]
                        rows = list(sheet.iter_rows(values_only=True))
                        if v == "VOL 5" and test_order == 8 and p_num == 1:
                            rows = rows[:180]
                        
                        for r in rows:
                            if not r or r[0] is None:
                                continue
                            w_text = str(r[0]).strip()
                            if not w_text:
                                continue
                            m_text = str(r[1]).strip() if len(r) > 1 and r[1] is not None else ""
                            words_raw.append((w_text, m_text))
                            
                seen = set()
                order = 1
                for w_text, m_text in words_raw:
                    k = w_text.lower()
                    if k not in seen:
                        seen.add(k)
                        word_id = str(uuid.uuid5(uuid.UUID(passage_id), f"word.{order}.{k}"))
                        words.append({
                            "id": word_id,
                            "topic_id": topic_id,
                            "passage_id": passage_id,
                            "word": w_text,
                            "meaning": m_text,
                            "lesson_name": test_name,
                            "lesson_order": test_order - 1,
                            "word_order": order
                        })
                        order += 1
                        
    return topics, tests, passages, words

def upload_data(topics, tests, passages, words, dry_run=False):
    print("\n--- Summary of Extracted Data ---")
    print(f"Topics:   {len(topics)}")
    print(f"Tests:    {len(tests)}")
    print(f"Passages: {len(passages)}")
    print(f"Words:    {len(words)}")
    
    if dry_run:
        print("Dry run requested. Skipping upload.")
        return
        
    headers = {
        "apikey": ANON_KEY,
        "Authorization": f"Bearer {ANON_KEY}",
        "Content-Type": "application/json"
    }
    
    # 1. Upload metadata
    print("\nUploading topics, tests, passages...")
    meta_url = f"{SUPABASE_URL}/rest/v1/rpc/bulk_insert_vol_metadata"
    meta_payload = {
        "p_topics": topics,
        "p_tests": tests,
        "p_passages": passages
    }
    resp = requests.post(meta_url, headers=headers, json=meta_payload, timeout=60)
    if not resp.ok:
        raise RuntimeError(f"Failed to upload metadata: HTTP {resp.status_code}: {resp.text}")
    print("Metadata uploaded successfully!")
    
    # 2. Upload words in batches
    words_url = f"{SUPABASE_URL}/rest/v1/rpc/bulk_insert_vol_words"
    batch_size = 1500
    total_words = len(words)
    total_inserted = 0
    
    print(f"\nUploading {total_words} words in batches of {batch_size}...")
    for i in range(0, total_words, batch_size):
        chunk = words[i:i + batch_size]
        batch_num = (i // batch_size) + 1
        total_batches = (total_words + batch_size - 1) // batch_size
        
        resp = requests.post(words_url, headers=headers, json={"p_words": chunk}, timeout=60)
        if not resp.ok:
            raise RuntimeError(f"Failed batch {batch_num}/{total_batches}: HTTP {resp.status_code}: {resp.text}")
        
        inserted = resp.json()
        total_inserted += (inserted if isinstance(inserted, int) else len(chunk))
        pct = min(100.0, ((i + len(chunk)) / total_words) * 100)
        print(f"  Batch {batch_num:2d}/{total_batches}: Uploaded {len(chunk):4d} words ({pct:5.1f}%)")
        
    print(f"\nCompleted! Total words processed: {total_words}, inserted: {total_inserted}")

if __name__ == "__main__":
    dry_run = "--dry-run" in sys.argv
    topics, tests, passages, words = parse_all_data()
    upload_data(topics, tests, passages, words, dry_run=dry_run)
